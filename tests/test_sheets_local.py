from openpyxl import load_workbook

from jobbank.models import Job
from jobbank.sheets import JOB_BANK, JOB_BANK_COLS, LocalXlsxBackend


def make_job(i, score=70):
    job = Job(source="t", company=f"Co{i}", title="SWE Placement 2027",
              url=f"https://x.com/{i}")
    job.score, job.category, job.score_reason = score, "SWE", "test"
    return job


def test_layout_and_append(tmp_path):
    path = tmp_path / "sheet.xlsx"
    be = LocalXlsxBackend(path)
    be.ensure_layout([{"name": "McLaren Racing", "ats": "careers", "url": "https://x"}])
    be.append_job_bank([make_job(1), make_job(2)])
    wb = load_workbook(path)
    assert set([JOB_BANK, "Opening Soon", "Archive", "Stats", "Watchlist",
                "Paste-in", "Run Log"]) <= set(wb.sheetnames)
    ws = wb[JOB_BANK]
    assert [c.value for c in ws[1]] == JOB_BANK_COLS
    assert ws.max_row == 3


def test_status_never_overwritten(tmp_path):
    path = tmp_path / "sheet.xlsx"
    be = LocalXlsxBackend(path)
    be.ensure_layout([])
    be.append_job_bank([make_job(1)])

    # Aug edits Status + Notes by hand
    wb = load_workbook(path)
    ws = wb[JOB_BANK]
    status_col = JOB_BANK_COLS.index("Status") + 1
    notes_col = JOB_BANK_COLS.index("Notes") + 1
    ws.cell(2, status_col, "Interview")
    ws.cell(2, notes_col, "spoke to recruiter")
    wb.save(path)

    # next weekly run appends more rows
    be2 = LocalXlsxBackend(path)
    assert "https://x.com/1" in be2.existing_links()
    be2.append_job_bank([make_job(2)])

    ws = load_workbook(path)[JOB_BANK]
    assert ws.cell(2, status_col).value == "Interview"
    assert ws.cell(2, notes_col).value == "spoke to recruiter"
    assert ws.cell(3, status_col).value == "Not applied"
